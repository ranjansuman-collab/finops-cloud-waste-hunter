import docx
import win32com.client
import os
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

abs_docx = os.path.abspath('Cloud_Waste_Hunter_Project_Report.docx')
abs_pdf = os.path.abspath('Cloud_Waste_Hunter_Project_Report.pdf')

doc_full = docx.Document(abs_docx)

for s in doc_full.sections:
    # Header
    hdr = s.header
    p_hdr = hdr.paragraphs[0]
    p_hdr.text = ''
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run('Cloud Waste Hunter — Technical Project Report | ISWE406')
    r_hdr.font.name = 'Times New Roman'
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.italic = True
    r_hdr.font.color.rgb = RGBColor(148, 163, 184)
    
    # Footer
    ftr = s.footer
    p_ftr = ftr.paragraphs[0]
    p_ftr.text = ''
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ftr_lbl = p_ftr.add_run('Page ')
    r_ftr_lbl.font.name = 'Times New Roman'
    r_ftr_lbl.font.size = Pt(9.5)
    r_ftr_lbl.font.color.rgb = RGBColor(100, 116, 139)
    
    r_num = p_ftr.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r_num._r.append(fldChar1)
    r_num._r.append(instrText)
    r_num._r.append(fldChar2)
    r_num._r.append(fldChar3)
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(9.5)
    r_num.font.color.rgb = RGBColor(100, 116, 139)

doc_full.save(abs_docx)
print("Updated and saved DOCX cleanly.")

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0

print("Opening in Word...")
doc = word.Documents.Open(abs_docx, False, True)

print("Computing statistics...")
pages = doc.ComputeStatistics(2)
words = doc.ComputeStatistics(0)
print(f"Page Count: {pages}")
print(f"Word Count: {words}")

print("Exporting PDF...")
doc.SaveAs(abs_pdf, FileFormat=17) # wdFormatPDF
print(f"Exported PDF to: {abs_pdf}")

doc.Close(0)
word.Quit()
print("ALL COMPLETE!")
