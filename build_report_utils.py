"""
build_report.py
Constructs the comprehensive, professional academic project report for Cloud Waste Hunter.
Outputs:
  - Cloud_Waste_Hunter_Project_Report.docx
  - Cloud_Waste_Hunter_Project_Report.pdf
Conforms strictly to academic report standards (A4, Times New Roman, 1.5 line spacing, 1-inch margins).
"""

import os
import sys
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- Color Palette Constants ---
PRIMARY_DARK = "0f172a"      # Slate 900
PRIMARY_BLUE = "0284c7"      # Sky 600
TEXT_MAIN = "1e293b"         # Slate 800
TEXT_MUTED = "64748b"        # Slate 500
BORDER_COLOR = "cbd5e1"      # Slate 300
BG_LIGHT = "f8fafc"          # Slate 50
TABLE_HEADER_BG = "1e293b"   # Slate 800
TABLE_HEADER_TEXT = "ffffff" # White
ROW_ALT_BG = "f1f5f9"        # Slate 100
ACCENT_GREEN = "16a34a"      # Green 600
ACCENT_RED = "dc2626"        # Red 600

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="cbd5e1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_page_number_to_paragraph(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(100, 116, 139)
    
    r_num = paragraph.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r_num._r.append(fldChar1)
    r_num._r.append(instrText)
    r_num._r.append(fldChar2)
    r_num._r.append(fldChar3)
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(9.5)
    r_num.font.color.rgb = RGBColor(100, 116, 139)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        
        # Header
        hdr = s.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.text = ""
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Cloud Waste Hunter — Technical Project Report | ISWE406")
        r_hdr.font.name = "Times New Roman"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.italic = True
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        ftr = s.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.text = ""
        add_page_number_to_paragraph(p_ftr)

print("Base setup utilities loaded successfully.")
